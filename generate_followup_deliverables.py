from __future__ import annotations

import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas as pdfcanvas


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Livrables_suivi_23-10-2026"
OUT.mkdir(exist_ok=True)
RESULTS = ROOT / "results"
DIAG = json.loads((RESULTS / "diagnostics.json").read_text(encoding="utf-8"))

NAVY = "153A5B"
CYAN = "168AAD"
ORANGE = "E07A3F"
LIGHT = "EAF2F6"
DARK = "24323D"
MUTED = "60717D"

pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))


def hex_color(value):
    return colors.HexColor("#" + value)


def report_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleX", parent=base["Title"], fontName="Arial-Bold", fontSize=22, leading=27, textColor=hex_color(NAVY), alignment=TA_CENTER, spaceAfter=12),
        "subtitle": ParagraphStyle("SubtitleX", parent=base["Normal"], fontName="Arial", fontSize=12, leading=17, textColor=hex_color(MUTED), alignment=TA_CENTER, spaceAfter=14),
        "h1": ParagraphStyle("H1X", parent=base["Heading1"], fontName="Arial-Bold", fontSize=15, leading=19, textColor=hex_color(NAVY), spaceBefore=12, spaceAfter=7),
        "h2": ParagraphStyle("H2X", parent=base["Heading2"], fontName="Arial-Bold", fontSize=11.5, leading=15, textColor=hex_color(CYAN), spaceBefore=9, spaceAfter=5),
        "body": ParagraphStyle("BodyX", parent=base["BodyText"], fontName="Arial", fontSize=9.3, leading=13.2, textColor=hex_color(DARK), alignment=TA_LEFT, spaceAfter=6),
        "bullet": ParagraphStyle("BulletX", parent=base["BodyText"], fontName="Arial", fontSize=9.1, leading=12.8, textColor=hex_color(DARK), leftIndent=15, firstLineIndent=-8, spaceAfter=3),
        "quote": ParagraphStyle("QuoteX", parent=base["BodyText"], fontName="Arial", fontSize=9.4, leading=13.3, textColor=hex_color(NAVY), borderColor=hex_color(CYAN), borderWidth=1.2, borderPadding=8, backColor=hex_color(LIGHT), spaceBefore=5, spaceAfter=8),
        "small": ParagraphStyle("SmallX", parent=base["BodyText"], fontName="Arial", fontSize=7.6, leading=10, textColor=hex_color(MUTED)),
    }


ST = report_styles()


def report_header_footer(canvas, doc):
    canvas.saveState()
    w, h = A4
    canvas.setFillColor(hex_color(NAVY))
    canvas.rect(0, h - 0.48 * cm, w, 0.48 * cm, fill=1, stroke=0)
    canvas.setFont("Arial", 7.5)
    canvas.setFillColor(hex_color(MUTED))
    canvas.drawString(2.1 * cm, 1.05 * cm, "Projet PI² n°176 — Rapport d’avancement n°1")
    canvas.drawRightString(w - 2.1 * cm, 1.05 * cm, f"Page {doc.page}")
    canvas.restoreState()


def build_report_pdf():
    path = OUT / "Rapport_avancement_20-10-2026.pdf"
    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=2.1 * cm, rightMargin=2.1 * cm, topMargin=1.7 * cm, bottomMargin=1.7 * cm, title="Rapport d'avancement PI2 n°176")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates(PageTemplate(id="main", frames=frame, onPage=report_header_footer))
    p, h1, h2 = ST["body"], ST["h1"], ST["h2"]
    story = [
        Spacer(1, 2.4 * cm),
        Paragraph("Projet PI² n°176", ST["title"]),
        Paragraph("Dynamic of RND and Forecasting the Price of Derivatives", ParagraphStyle("BigSub", parent=ST["h1"], alignment=TA_CENTER, fontSize=17, leading=22)),
        Spacer(1, 0.8 * cm),
        Paragraph("Rapport d’avancement n°1", ParagraphStyle("MidSub", parent=ST["subtitle"], fontSize=14, textColor=hex_color(CYAN))),
        Paragraph("20 octobre 2026", ST["subtitle"]),
        Spacer(1, 1.5 * cm),
        Paragraph("Rayan Ait Arkoub · Matthieu Balliste · Ilan Chahdi<br/>Nathan Denoux · Jordi Paris", ParagraphStyle("Authors", parent=p, alignment=TA_CENTER, fontSize=11, leading=17)),
        Spacer(1, 2.1 * cm),
        Paragraph("Tuteur : Matthieu Garcin", ParagraphStyle("Tutor", parent=p, alignment=TA_CENTER, textColor=hex_color(MUTED))),
        PageBreak(),
        Paragraph("Résumé exécutif", h1),
        Paragraph(
            "Le projet vise à construire des séries temporelles de distributions risque-neutre à partir d’options, à comparer trois méthodes d’estimation, à modéliser leur dynamique et à tester une application à la couverture. Ce premier jalon fixe le périmètre SPX, documente les accès aux données et livre un prototype Breeden–Litzenberger reproductible en Python et Julia.", p),
        Paragraph(
            f"Une photographie publique Cboe horodatée au {DIAG['quote_timestamp']} contient {DIAG['selected_contracts_before_filtering']:,} contrats pour l’échéance du 15 janvier 2027. Après nettoyage, {DIAG['valid_contract_quotes']:,} cotations sont valides et {DIAG['calibration_points']} points alimentent une calibration SVI à cinq paramètres. La masse brute de la densité sur la grille atteint {DIAG['raw_density_mass_on_grid']:.4f}.", ST["quote"]),
        Paragraph("1. Problématique et objectif", h1),
        Paragraph("Une volatilité Black–Scholes unique ne permet pas de reproduire simultanément tous les prix d’options d’une même échéance. Une volatilité implicite peut ajuster chaque contrat séparément, mais le smile ainsi obtenu révèle qu’un modèle gaussien à paramètre unique est insuffisant. Le projet cherche donc une distribution risque-neutre commune à l’ensemble des strikes.", p),
        Paragraph("La chaîne scientifique retenue est : données d’options → estimation de trois RND → stabilité temporelle → réduction factorielle → prévision → couverture ou réplication.", p),
        Paragraph("2. Cadre théorique", h1),
        Paragraph("Pour une option européenne, le prix du call est l’espérance actualisée de son payoff sous la mesure risque-neutre. Sous des conditions de régularité, Breeden–Litzenberger établit que la densité est la dérivée seconde du prix du call par rapport au strike, divisée par le facteur d’actualisation.", p),
        Paragraph("q(K,T) = [1 / D(t,T)] × ∂²C(K,T) / ∂K²", ST["quote"]),
        Paragraph("La double dérivation amplifie fortement le bruit. Le prototype l’applique donc à une courbe de calls lissée, et non directement aux derniers prix de transaction.", p),
        Paragraph("3. Accès aux données", h1),
        Paragraph("L’accès à la chaîne SPX publique et différée de Cboe est confirmé. Bloomberg n’est pas accessible depuis cette machine : aucun terminal ni connecteur API n’est installé. Une extraction historique devra être réalisée depuis un poste Bloomberg de l’école ou obtenue via OptionMetrics/Cboe DataShop. L’API publique permet cependant de valider la totalité du pipeline sur une coupe réelle.", p),
    ]

    rows = [
        ["Indicateur", "Valeur"],
        ["Horodatage", DIAG["quote_timestamp"]],
        ["Spot SPX", f"{DIAG['spot']:,.2f}"],
        ["Échéance", "15 janvier 2027"],
        ["Maturité", f"{DIAG['maturity_days']} jours"],
        ["Contrats avant filtres", f"{DIAG['selected_contracts_before_filtering']:,}"],
        ["Cotations valides", f"{DIAG['valid_contract_quotes']:,}"],
        ["Points de calibration", f"{DIAG['calibration_points']:,}"],
    ]
    tbl = Table(rows, colWidths=[7.5 * cm, 7.5 * cm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), hex_color(NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Arial-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Arial"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.7), ("GRID", (0, 0), (-1, -1), 0.35, hex_color("B9C8D1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hex_color(LIGHT)]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([tbl, Spacer(1, 8), Image(str(RESULTS / "01_data_quality.png"), width=16.1 * cm, height=6.15 * cm), Paragraph("Figure 1 — Spreads relatifs et ancienneté des derniers échanges.", ST["small"]),
        Paragraph("4. Nettoyage et calibration", h1),
        Paragraph("Le nettoyage privilégie le mid bid-ask. Les bids nuls, marchés croisés, spreads relatifs supérieurs à 35 % et strikes excessivement éloignés sont exclus. Le forward et le facteur d’actualisation sont estimés par une régression pondérée de la parité call-put C − P = D(F − K) autour de la monnaie.", p),
        Paragraph("Les options hors de la monnaie sont combinées par parité pour obtenir une courbe de calls. Les volatilités sont ensuite ajustées par une SVI brute à cinq paramètres, ce qui répond à l’objectif d’un modèle commun parcimonieux de la coupe de strikes.", p),
        Image(str(RESULTS / "02_svi_smile.png"), width=14.9 * cm, height=8.4 * cm), Paragraph("Figure 2 — Smile observée et calibration SVI à cinq paramètres.", ST["small"]),
        Paragraph("5. Résultats du prototype", h1),
        Paragraph(f"La parité fournit un forward de {DIAG['forward_from_parity']:,.2f}, un facteur d’actualisation de {DIAG['discount_factor_from_parity']:.5f} et un taux continu implicite de {100*DIAG['implied_continuous_rate']:.2f} %. La RMSE de parité est de {DIAG['put_call_parity_rmse']:.3f} point et la RMSE de reconstruction des calls de {DIAG['call_price_rmse']:.3f} point.", p),
        Image(str(RESULTS / "03_call_curve.png"), width=14.9 * cm, height=8.4 * cm), Paragraph("Figure 3 — Prix utilisés et courbe de calls reconstruite.", ST["small"]),
        Paragraph(f"La masse brute obtenue sur la grille vaut {DIAG['raw_density_mass_on_grid']:.5f}. La masse négative absolue est limitée à {DIAG['absolute_negative_mass']:.6f}. Bien que {100*DIAG['negative_grid_fraction']:.1f} % des points de l’extrême queue soient légèrement négatifs, aucun ne dépasse en amplitude 0,1 % du pic de densité. Après troncature des seules valeurs négatives et normalisation, la moyenne vaut {DIAG['normalized_density_mean']:,.2f}.", p),
        Image(str(RESULTS / "04_risk_neutral_density.png"), width=14.9 * cm, height=8.4 * cm), Paragraph("Figure 4 — Première RND SPX par Breeden–Litzenberger.", ST["small"]),
        Paragraph("6. Validation Julia", h1),
        Paragraph("Le calcul de la dérivée seconde, de la masse, de la normalisation et des moments a été réimplémenté en Julia 1.10.10 sans package externe. L’écart absolu moyen entre les densités normalisées Python et Julia est de 3,3 × 10^-9 ; l’écart maximal de 2,1 × 10^-6 se situe aux bords, où les conventions de différences finies varient.", p),
        Paragraph("7. Limites", h1),
    ])
    for item in [
        "La photographie publique Cboe ne remplace pas un historique quotidien sur un mois.",
        "Le seuil de spread de 35 % est provisoire et doit faire l’objet d’une analyse de sensibilité.",
        "La SVI brute n’impose pas toutes les contraintes d’absence d’arbitrage butterfly ; la faible négativité de queue doit être corrigée par eSSVI ou un lissage convexe.",
        "Une maturité cible constante est nécessaire pour séparer dynamique de marché et diminution mécanique de l’échéance.",
        "Les performances de prévision et de couverture devront être mesurées hors échantillon et nettes de coûts.",
    ]:
        story.append(Paragraph("• " + item, ST["bullet"]))
    story.extend([
        Paragraph("8. Prochaines étapes", h1),
        Paragraph("Après obtention de l’historique, les méthodes Breeden–Litzenberger, Fourier-cosinus et distorsion de la probabilité historique seront comparées sur une grille commune de log-moneyness. Les densités seront étudiées avec KL, Jensen–Shannon et Wasserstein. Une réduction factorielle permettra de prévoir les scores par AR/VAR, puis de reconstruire la densité future et son incertitude.", p),
        Paragraph("L’application finale comparera une couverture fondée sur la RND courante, une couverture utilisant la RND prévue et un benchmark Black–Scholes, selon l’erreur de couverture et les coûts de transaction.", p),
        Paragraph("9. Conclusion", h1),
        Paragraph("Le premier jalon est atteint : accès Cboe courant confirmé, échantillon SPX réel extrait, pipeline Breeden–Litzenberger opérationnel, calibration commune à cinq paramètres, diagnostics produits et validation Julia effectuée. Le principal risque résiduel est désormais externe : l’accès à une profondeur historique suffisante.", ST["quote"]),
        Paragraph("Références principales", h1),
        Paragraph("Breeden, D. T. & Litzenberger, R. H. (1978). Prices of State-Contingent Claims Implicit in Option Prices. The Journal of Business, 51(4), 621–651.", p),
        Paragraph("Figlewski, S. (2010). Estimating the Implied Risk-Neutral Density for the US Market Portfolio. Oxford University Press, 323–353.", p),
        Paragraph("Figlewski, S. (2018). Risk-Neutral Densities: A Review. Annual Review of Financial Economics, 10, 329–359.", p),
        Paragraph("Cui, Z. & Yu, Z. (2021). A Model-Free Fourier Cosine Method for Estimating the Risk-Neutral Density. The Journal of Derivatives, 29(2), 149–171.", p),
        Paragraph("Giacomini, E., Härdle, W. K. & Krätschmer, V. (2009). Dynamic Semiparametric Factor Models in Risk Neutral Density Estimation. AStA, 93(4), 387–402.", p),
        Paragraph("Taboga, M. (2016). Option-Implied Probability Distributions: How Reliable? How Jagged? International Review of Economics & Finance, 45, 453–469.", p),
    ])
    doc.build(story)
    return path


def set_background(slide, prs, color="FFFFFF"):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_slide_header(slide, title, number, prs):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(NAVY); bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.38), Inches(11.9), Inches(0.62))
    p = box.text_frame.paragraphs[0]; p.text = title; p.font.name = "Arial"; p.font.bold = True; p.font.size = Pt(25); p.font.color.rgb = RGBColor.from_string(NAVY)
    foot = slide.shapes.add_textbox(Inches(0.65), Inches(7.12), Inches(12), Inches(0.2))
    fp = foot.text_frame.paragraphs[0]; fp.text = f"Projet PI² n°176 — Suivi n°1     {number}/10"; fp.font.name = "Arial"; fp.font.size = Pt(8); fp.font.color.rgb = RGBColor.from_string(MUTED)


def add_bullets(slide, items, x, y, w, h, font_size=18, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.03)
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = item; para.level = 0; para.font.name = "Arial"; para.font.size = Pt(font_size); para.font.color.rgb = RGBColor.from_string(color); para.space_after = Pt(10)
        para.text = "• " + para.text
    return box


def add_metric(slide, value, label, x, y, w=2.7):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.15))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(LIGHT); shape.line.color.rgb = RGBColor.from_string(CYAN)
    tf = shape.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]; p1.text = value; p1.alignment = PP_ALIGN.CENTER; p1.font.name = "Arial"; p1.font.bold = True; p1.font.size = Pt(24); p1.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = tf.add_paragraph(); p2.text = label; p2.alignment = PP_ALIGN.CENTER; p2.font.name = "Arial"; p2.font.size = Pt(10); p2.font.color.rgb = RGBColor.from_string(MUTED)


def add_picture(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides = []
    slide = prs.slides.add_slide(blank); set_background(slide, prs)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.25)); bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(NAVY); bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.7), Inches(1.8)); tf=box.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text="Projet PI² n°176"; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(30); p.font.color.rgb=RGBColor.from_string(NAVY); p.alignment=PP_ALIGN.CENTER
    p=tf.add_paragraph(); p.text="Dynamic of RND and Forecasting the Price of Derivatives"; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(23); p.font.color.rgb=RGBColor.from_string(NAVY); p.alignment=PP_ALIGN.CENTER; p.space_before=Pt(10)
    p=tf.add_paragraph(); p.text="Réunion de suivi n°1 — 23 octobre 2026"; p.font.name="Arial"; p.font.size=Pt(17); p.font.color.rgb=RGBColor.from_string(CYAN); p.alignment=PP_ALIGN.CENTER; p.space_before=Pt(18)
    box2=slide.shapes.add_textbox(Inches(1.2), Inches(4.45), Inches(10.9), Inches(1)); p=box2.text_frame.paragraphs[0]; p.text="Rayan Ait Arkoub · Matthieu Balliste · Ilan Chahdi · Nathan Denoux · Jordi Paris"; p.font.name="Arial"; p.font.size=Pt(16); p.font.color.rgb=RGBColor.from_string(DARK); p.alignment=PP_ALIGN.CENTER
    slides.append(("Titre", [], None))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"1. Problématique",2,prs)
    add_bullets(slide,["Une volatilité Black–Scholes unique ne reproduit pas toute la chaîne d’options.","Chaque volatilité implicite ajuste un contrat, mais le smile exige une distribution commune.","Objectif : estimer, suivre et prévoir la RND, puis l’utiliser pour la couverture."],0.8,1.35,5.7,4.7,20)
    # pipeline
    labels=["Options","RND","Stabilité","Prévision","Couverture"]
    for i,label in enumerate(labels):
        x=6.7+i*1.23; shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.25),Inches(1.05),Inches(1.0)); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(LIGHT); shape.line.color.rgb=RGBColor.from_string(CYAN); tf=shape.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=label; p.alignment=PP_ALIGN.CENTER; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(11); p.font.color.rgb=RGBColor.from_string(NAVY)
        if i<4:
            arr=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+1.05),Inches(2.55),Inches(0.22),Inches(0.35)); arr.fill.solid(); arr.fill.fore_color.rgb=RGBColor.from_string(ORANGE); arr.line.fill.background()
    slides.append(("Problématique",["Une volatilité Black–Scholes unique ne reproduit pas toute la chaîne d’options.","Chaque volatilité implicite ajuste un contrat, mais le smile exige une distribution commune.","Objectif : estimer, suivre et prévoir la RND, puis l’utiliser pour la couverture."],None))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"2. Données et accès",3,prs)
    add_metric(slide,f"{DIAG['selected_contracts_before_filtering']:,}","contrats SPX",0.8,1.35)
    add_metric(slide,f"{DIAG['valid_contract_quotes']:,}","cotations valides",3.8,1.35)
    add_metric(slide,f"{DIAG['calibration_points']}","points de calibration",6.8,1.35)
    add_metric(slide,"91 jours","échéance 15/01/2027",9.8,1.35)
    add_picture(slide,RESULTS/"01_data_quality.png",0.85,2.85,7.2,3.6)
    add_bullets(slide,["Cboe public différé : accès confirmé.","Bloomberg : terminal/API absent ici.","Historique : export école, OptionMetrics ou DataShop requis.","Mid bid-ask préféré au last price souvent ancien."],8.35,3.0,4.15,3.3,15)
    slides.append(("Données et accès",[f"{DIAG['selected_contracts_before_filtering']:,} contrats ; {DIAG['valid_contract_quotes']:,} cotations valides ; {DIAG['calibration_points']} points calibrés.","Cboe public différé : accès confirmé.","Bloomberg : terminal/API absent sur cette machine.","Historique : export école, OptionMetrics ou DataShop requis."],RESULTS/"01_data_quality.png"))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"3. Pipeline de nettoyage",4,prs)
    steps=[("1","Sélection échéance ≈ 3 mois"),("2","Contrôles bid–ask et spreads"),("3","Parité call-put : F et D"),("4","OTM calls + puts synthétiques"),("5","Grille commune log(K/F)")]
    for i,(n,label) in enumerate(steps):
        y=1.15+i*1.05; circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(0.9),Inches(y),Inches(0.62),Inches(0.62)); circle.fill.solid(); circle.fill.fore_color.rgb=RGBColor.from_string(CYAN); circle.line.fill.background(); tf=circle.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text=n; p.alignment=PP_ALIGN.CENTER; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(18); p.font.color.rgb=RGBColor(255,255,255)
        box=slide.shapes.add_textbox(Inches(1.75),Inches(y-0.02),Inches(5),Inches(0.7)); p=box.text_frame.paragraphs[0]; p.text=label; p.font.name="Arial"; p.font.size=Pt(19); p.font.color.rgb=RGBColor.from_string(DARK)
    add_bullets(slide,["Exclus : bid nul, marché croisé, spread relatif > 35 %.","Volume et open interest conservés comme diagnostics.","Tous les seuils seront testés en sensibilité."],7.1,1.65,5.1,3.7,18)
    slides.append(("Pipeline de nettoyage",[x[1] for x in steps]+["Exclus : bid nul, marché croisé et spread relatif supérieur à 35 %.","Seuils à tester en analyse de sensibilité."],None))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"4. Calibration commune : SVI",5,prs)
    add_picture(slide,RESULTS/"02_svi_smile.png",0.75,1.25,7.5,4.9)
    svi=DIAG["svi_parameters"]
    add_bullets(slide,["5 paramètres pour toute la coupe de strikes.",f"Forward estimé : {DIAG['forward_from_parity']:,.1f}",f"Taux continu implicite : {100*DIAG['implied_continuous_rate']:.2f} %",f"RMSE des calls : {DIAG['call_price_rmse']:.3f} point",f"ρ = {svi['rho']:.3f} ; σ = {svi['sigma']:.3f}"],8.55,1.55,3.95,4.5,17)
    slides.append(("Calibration SVI",["Une seule smile à cinq paramètres explique 338 points.",f"Forward estimé : {DIAG['forward_from_parity']:,.1f}.",f"Taux continu implicite : {100*DIAG['implied_continuous_rate']:.2f} %.",f"RMSE des calls : {DIAG['call_price_rmse']:.3f} point."],RESULTS/"02_svi_smile.png"))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"5. Breeden–Litzenberger",6,prs)
    add_picture(slide,RESULTS/"03_call_curve.png",0.75,1.25,7.45,4.9)
    formula=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(8.5),Inches(1.45),Inches(4.0),Inches(1.05)); formula.fill.solid(); formula.fill.fore_color.rgb=RGBColor.from_string(LIGHT); formula.line.color.rgb=RGBColor.from_string(CYAN); tf=formula.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text="q(K,T) = (1/D) × ∂²C/∂K²"; p.alignment=PP_ALIGN.CENTER; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(19); p.font.color.rgb=RGBColor.from_string(NAVY)
    add_bullets(slide,["Dérivation de la courbe lissée, pas des prix bruts.","Grille régulière de 1 601 strikes.","Contrôles : masse, positivité, moyenne et convexité."],8.55,2.9,3.8,3.0,17)
    slides.append(("Breeden–Litzenberger",["q(K,T) = (1/D) × ∂²C/∂K².","Dérivation de la courbe lissée, pas des prix bruts.","Grille régulière de 1 601 strikes.","Contrôles : masse, positivité, moyenne et convexité."],RESULTS/"03_call_curve.png"))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"6. Première densité risque-neutre",7,prs)
    add_picture(slide,RESULTS/"04_risk_neutral_density.png",0.7,1.2,7.6,5.05)
    add_metric(slide,f"{DIAG['raw_density_mass_on_grid']:.4f}","masse brute",8.6,1.45,3.4)
    add_metric(slide,f"{DIAG['normalized_density_mean']:,.0f}","moyenne normalisée",8.6,2.9,3.4)
    add_metric(slide,f"{DIAG['absolute_negative_mass']:.6f}","masse négative",8.6,4.35,3.4)
    box=slide.shapes.add_textbox(Inches(8.55),Inches(5.85),Inches(3.9),Inches(0.65)); p=box.text_frame.paragraphs[0]; p.text="Aucune négativité matérielle > 0,1 % du pic"; p.font.name="Arial"; p.font.size=Pt(13); p.font.color.rgb=RGBColor.from_string(ORANGE); p.alignment=PP_ALIGN.CENTER
    slides.append(("Première RND",[f"Masse brute : {DIAG['raw_density_mass_on_grid']:.4f}.",f"Moyenne normalisée : {DIAG['normalized_density_mean']:,.0f}.",f"Masse négative : {DIAG['absolute_negative_mass']:.6f}.","Aucune négativité matérielle supérieure à 0,1 % du pic."],RESULTS/"04_risk_neutral_density.png"))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"7. Structure par terme des RND",8,prs)
    add_picture(slide,RESULTS/"05_rnd_surface_3d_maturity.png",0.55,1.08,8.15,5.75)
    add_bullets(slide,["Six maturités : 31 à 182 jours.","Grille commune de log-moneyness.","La densité s’élargit avec l’horizon.","Il s’agit d’une structure par terme à date fixe, pas encore d’une dynamique temporelle."],8.95,1.45,3.75,4.7,16)
    slides.append(("Structure par terme des RND",["Six maturités : 31, 63, 91, 126, 154 et 182 jours.","Grille commune de log-moneyness.","La densité s’élargit avec l’horizon.","Surface à date fixe : l’axe vertical est la maturité, pas le temps calendaire."],RESULTS/"05_rnd_surface_3d_maturity.png"))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"8. Validation et limites",9,prs)
    add_bullets(slide,["Calcul BL réimplémenté en Julia 1.10.10 sans package externe.","Écart moyen Python–Julia : 3,3 × 10^-9.","Écart maximal aux bords : 2,1 × 10^-6."],0.85,1.3,5.7,2.8,19)
    add_bullets(slide,["Pas encore d’historique quotidien.","Seuil de spread à tester.","SVI brute pas totalement sans arbitrage butterfly.","Besoin d’une maturité constante pour la dynamique."],7.0,1.3,5.2,3.8,18,color=ORANGE)
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(1.1),Inches(4.65),Inches(10.9),Inches(1.15)); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(LIGHT); shape.line.color.rgb=RGBColor.from_string(CYAN); tf=shape.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text="Conclusion technique : pipeline reproductible validé ; profondeur historique encore à sécuriser."; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(20); p.font.color.rgb=RGBColor.from_string(NAVY); p.alignment=PP_ALIGN.CENTER
    slides.append(("Validation et limites",["Calcul BL réimplémenté en Julia sans package externe.","Écart moyen Python–Julia : 3,3 × 10^-9.","Pas encore d’historique quotidien.","Seuil de spread provisoire.","SVI brute pas totalement sans arbitrage butterfly.","Besoin d’une maturité constante pour la dynamique."],None))

    slide=prs.slides.add_slide(blank); set_background(slide, prs); add_slide_header(slide,"9. Prochaines étapes et décisions attendues",10,prs)
    add_bullets(slide,["Obtenir un historique SPX quotidien à maturité ≈ 3 mois.","Comparer BL, Fourier-cosinus et distorsion historique.","Mesurer stabilité : KL, Jensen–Shannon, Wasserstein.","Réduire les RND en facteurs, puis prévoir par AR/VAR.","Comparer les erreurs de couverture nettes de coûts."],0.85,1.25,6.0,4.8,18)
    add_bullets(slide,["Accès historique à utiliser ?","Maturité constante validée ?","Trois méthodes intégrales attendues ?","Couverture prioritaire au trading ?"],7.3,1.6,4.9,3.5,19,color=NAVY)
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(7.35),Inches(5.15),Inches(4.7),Inches(0.95)); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(ORANGE); shape.line.fill.background(); tf=shape.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.text="Objectif suivant : une série temporelle de RND comparables"; p.font.name="Arial"; p.font.bold=True; p.font.size=Pt(17); p.font.color.rgb=RGBColor(255,255,255); p.alignment=PP_ALIGN.CENTER
    slides.append(("Prochaines étapes",["Obtenir un historique SPX quotidien à maturité proche de trois mois.","Comparer BL, Fourier-cosinus et distorsion historique.","Mesurer la stabilité par KL, Jensen–Shannon et Wasserstein.","Réduire les RND en facteurs et prévoir les scores par AR/VAR.","Comparer les erreurs de couverture nettes de coûts.","Décisions attendues : accès, maturité, méthodes obligatoires et priorité couverture/trading."],None))

    path = OUT / "Support_suivi_23-10-2026.pptx"
    prs.save(path)
    return path, slides


def build_slides_pdf(slides):
    path = OUT / "Support_suivi_23-10-2026.pdf"
    c = pdfcanvas.Canvas(str(path), pagesize=landscape(A4))
    W, H = landscape(A4)
    for i, (title, bullets, image_path) in enumerate(slides, start=1):
        c.setFillColor(colors.white); c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(hex_color(NAVY)); c.rect(0, H - 14, W, 14, fill=1, stroke=0)
        if i == 1:
            c.setFont("Arial-Bold", 26); c.setFillColor(hex_color(NAVY)); c.drawCentredString(W/2, H-125, "Projet PI² n°176")
            c.setFont("Arial-Bold", 20); c.drawCentredString(W/2, H-165, "Dynamic of RND and Forecasting the Price of Derivatives")
            c.setFont("Arial", 15); c.setFillColor(hex_color(CYAN)); c.drawCentredString(W/2, H-210, "Réunion de suivi n°1 — 23 octobre 2026")
            c.setFont("Arial", 12); c.setFillColor(hex_color(DARK)); c.drawCentredString(W/2, H-340, "Rayan Ait Arkoub · Matthieu Balliste · Ilan Chahdi · Nathan Denoux · Jordi Paris")
        else:
            c.setFont("Arial-Bold", 23); c.setFillColor(hex_color(NAVY)); c.drawString(38, H-52, title)
            if image_path:
                c.drawImage(str(image_path), 38, 75, width=470, height=285, preserveAspectRatio=True, anchor='c')
                x, y, width = 535, H-105, 270
            else:
                x, y, width = 65, H-115, 700
            c.setFillColor(hex_color(DARK)); c.setFont("Arial", 15)
            for bullet in bullets:
                words = bullet.split(); lines=[]; current=""
                for word in words:
                    candidate=(current+" "+word).strip()
                    if c.stringWidth(candidate,"Arial",15) > width and current:
                        lines.append(current); current=word
                    else: current=candidate
                if current: lines.append(current)
                c.setFillColor(hex_color(CYAN)); c.circle(x, y+4, 3, fill=1, stroke=0)
                c.setFillColor(hex_color(DARK))
                for line in lines:
                    c.drawString(x+14, y, line); y-=20
                y-=15
        c.setFont("Arial", 8); c.setFillColor(hex_color(MUTED)); c.drawString(38, 22, f"Projet PI² n°176 — Suivi n°1     {i}/10")
        c.showPage()
    c.save()
    return path


def copy_sources():
    for name in ["rapport_avancement_20-10-2026.tex", "references.bib", "run_prototype.ps1", "Notes_orales_suivi_23-10-2026.md", "requirements.txt", "README_prototype.md"]:
        shutil.copy2(ROOT / name, OUT / name)
    for name in ["spx_breeden_litzenberger.py", "breeden_litzenberger.jl"]:
        shutil.copy2(ROOT / "src" / name, OUT / name)


if __name__ == "__main__":
    report = build_report_pdf()
    pptx, slides = build_pptx()
    slides_pdf = build_slides_pdf(slides)
    copy_sources()
    manifest = OUT / "MANIFESTE.txt"
    manifest.write_text(
        "PI2 n°176 — Livrables du premier suivi\n\n"
        "À envoyer :\n- Rapport_avancement_20-10-2026.pdf\n- Support_suivi_23-10-2026.pdf ou .pptx\n\n"
        "Sources : rapport_avancement_20-10-2026.tex, references.bib, scripts Python/Julia.\n"
        "Le PDF est généré localement ; la source LaTeX se compile avec pdflatex/biber.\n",
        encoding="utf-8",
    )
    for path in sorted(OUT.iterdir()):
        print(f"{path.name}: {path.stat().st_size} bytes")
