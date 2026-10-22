"""Génère le support académique du suivi PI² n°176.

Le support privilégie les équations, les graphiques et les diagnostics quantitatifs.
Il est volontairement sobre afin de rester cohérent avec le rapport de recherche.
"""

from __future__ import annotations

import json
import html
import shutil
import subprocess
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from scipy.integrate import cumulative_trapezoid
from reportlab.lib import colors as rl_colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results"
ASSETS = RESULTS / "presentation"
OUT = ROOT / "Livrables_suivi_23-10-2026"
DIAG = json.loads((RESULTS / "diagnostics.json").read_text(encoding="utf-8"))
TERM = json.loads((RESULTS / "rnd_term_structure_diagnostics.json").read_text(encoding="utf-8"))

BLACK = "111111"
DARK = "27313A"
MUTED = "59636C"
BLUE = "315B78"
RED = "A33A32"
LIGHT = "E8ECEF"
WHITE = "FFFFFF"
TOTAL_SLIDES = 11


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def render_equation(filename: str, latex: str, width: float = 9.0, height: float = 0.75) -> Path:
    path = ASSETS / filename
    fig = plt.figure(figsize=(width, height), dpi=220, facecolor="white")
    fig.text(0.5, 0.5, f"${latex}$", ha="center", va="center", fontsize=22, color="#111111")
    fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.08, transparent=False)
    plt.close(fig)
    return path


def make_support_figures() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    equations = {
        "call": render_equation(
            "eq_call.png",
            r"C_t(K,T)=D(t,T)\,\mathbb{E}^{\mathbb{Q}}_t\!\left[(S_T-K)^+\right]",
        ),
        "parity": render_equation(
            "eq_parity.png",
            r"C_t(K,T)-P_t(K,T)=D(t,T)\left(F_{t,T}-K\right)",
        ),
        "bl": render_equation(
            "eq_bl.png",
            r"q_t(K,T)=\frac{1}{D(t,T)}\frac{\partial^2 C_t(K,T)}{\partial K^2}",
        ),
        "svi": render_equation(
            "eq_svi.png",
            r"w(k)=a+b\left[\rho(k-m)+\sqrt{(k-m)^2+\sigma^2}\right],\quad k=\log(K/F_{t,T})",
            width=10.0,
        ),
        "change": render_equation(
            "eq_change.png",
            r"q_k(k)=q_S\!\left(F_{t,T}e^k\right)F_{t,T}e^k",
        ),
        "factor": render_equation(
            "eq_factor.png",
            r"q_t(k)=\bar q(k)+\sum_{j=1}^{J}\beta_{j,t}\,\phi_j(k)+\varepsilon_t(k)",
        ),
        "var": render_equation(
            "eq_var.png",
            r"\boldsymbol{\beta}_t=\mathbf{c}+\sum_{\ell=1}^{p}\mathbf{A}_\ell\boldsymbol{\beta}_{t-\ell}+\boldsymbol{u}_t",
        ),
    }

    density = pd.read_csv(RESULTS / "rnd_python.csv")
    calibration = pd.read_csv(ROOT / "data" / "processed" / "spx_calibration_20261016_20270115.csv")
    strike = density["strike"].to_numpy()
    q = density["rnd_normalized"].to_numpy()
    cdf = cumulative_trapezoid(q, strike, initial=0.0)
    cdf /= cdf[-1]
    k_min, k_max = calibration["strike"].min(), calibration["strike"].max()
    forward = DIAG["forward_from_parity"]
    mean = DIAG["normalized_density_mean"]

    plt.style.use("seaborn-v0_8-whitegrid")
    fig, axes = plt.subplots(1, 2, figsize=(12.0, 4.25))
    axes[0].plot(strike, q, color="#315B78", lw=2.2)
    axes[0].axvspan(k_min, k_max, color="#B7C9D6", alpha=0.28, label="strikes utilisés")
    axes[0].axvline(forward, color="#111111", ls="--", lw=1.2, label=f"forward {forward:,.0f}")
    axes[0].axvline(mean, color="#A33A32", ls=":", lw=1.5, label=f"moyenne {mean:,.0f}")
    axes[0].set(xlabel="Niveau terminal du SPX", ylabel="Densité", title="Support observé et extrapolation")
    axes[0].legend(fontsize=8, loc="upper right")
    axes[1].plot(strike, cdf, color="#315B78", lw=2.2)
    axes[1].axvspan(k_min, k_max, color="#B7C9D6", alpha=0.28)
    axes[1].axhline(0.5, color="#777777", ls="--", lw=0.9)
    axes[1].axvline(forward, color="#111111", ls="--", lw=1.2)
    axes[1].set(xlabel="Niveau terminal du SPX", ylabel="Probabilité cumulée", title="Fonction de répartition implicite")
    fig.tight_layout()
    support_diag = ASSETS / "07_support_and_cdf.png"
    fig.savefig(support_diag, dpi=210, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    rows = [x for x in TERM["expiries"] if "error" not in x]
    days = np.array([x["maturity_days"] for x in rows])
    mass = np.array([x["raw_density_mass_on_grid"] for x in rows])
    mean_gap = np.array([
        100.0 * (x["normalized_density_mean"] - x["forward"]) / x["forward"] for x in rows
    ])
    colors = ["#A33A32" if d == 31 else "#315B78" for d in days]
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.1))
    axes[0].bar(days.astype(str), mass, color=colors, width=0.68)
    axes[0].axhline(1.0, color="#111111", ls="--", lw=1.1)
    axes[0].set_ylim(0.96, max(1.155, mass.max() + 0.01))
    axes[0].set(xlabel="Maturité (jours)", ylabel="Masse brute", title="Intégrale de la densité avant normalisation")
    axes[1].bar(days.astype(str), mean_gap, color=colors, width=0.68)
    axes[1].axhline(0.0, color="#111111", lw=1.0)
    axes[1].set(xlabel="Maturité (jours)", ylabel="Écart moyen–forward (%)", title="Condition de moment sous la mesure risque-neutre")
    fig.text(0.5, -0.01, "Rouge : coupe 31 jours non fiable en raison de l’extrapolation de l’aile droite SVI", ha="center", fontsize=9, color="#A33A32")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    term_diag = ASSETS / "08_term_structure_diagnostics.png"
    fig.savefig(term_diag, dpi=210, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    equations["support_diag"] = support_diag
    equations["term_diag"] = term_diag
    return equations


def set_background(slide, prs) -> None:
    shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_header(slide, title: str, number: int, prs) -> None:
    line = slide.shapes.add_shape(1, Inches(0.58), Inches(0.95), Inches(12.15), Inches(0.018))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(BLACK); line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.62), Inches(0.30), Inches(11.9), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = title; p.font.name = "Arial"; p.font.bold = True; p.font.size = Pt(25); p.font.color.rgb = rgb(BLACK)
    foot = slide.shapes.add_textbox(Inches(0.62), Inches(7.12), Inches(12.0), Inches(0.18))
    p = foot.text_frame.paragraphs[0]
    p.text = f"Projet PI² n°176  ·  Suivi n°1  ·  23 octobre 2026                                      {number}/{TOTAL_SLIDES}"
    p.font.name = "Arial"; p.font.size = Pt(8); p.font.color.rgb = rgb(MUTED)


def add_text(slide, text: str, x: float, y: float, w: float, h: float, size: int = 17,
             bold: bool = False, color: str = DARK, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = "Arial"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float,
                size: int = 17, color: str = DARK) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0.02); tf.margin_right = 0
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = "• " + item; p.font.name = "Arial"; p.font.size = Pt(size); p.font.color.rgb = rgb(color)
        p.space_after = Pt(10)


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(path) as image:
        ratio = image.width / image.height
    target_ratio = w / h
    if ratio >= target_ratio:
        width, height = w, w / ratio
        left, top = x, y + (h - height) / 2
    else:
        width, height = h * ratio, h
        left, top = x + (w - width) / 2, y
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def add_caption(slide, text: str, x: float, y: float, w: float) -> None:
    add_text(slide, text, x, y, w, 0.35, size=9, color=MUTED, align=PP_ALIGN.CENTER)


def add_metric_line(slide, value: str, label: str, x: float, y: float, w: float = 3.6) -> None:
    add_text(slide, value, x, y, 1.25, 0.38, size=20, bold=True, color=BLACK)
    add_text(slide, label, x + 1.3, y + 0.03, w - 1.3, 0.38, size=12, color=MUTED)


def build_presentation() -> Path:
    assets = make_support_figures()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — Titre
    slide = prs.slides.add_slide(blank); set_background(slide, prs)
    add_text(slide, "Projet PI² n°176", 0.8, 0.75, 11.7, 0.55, 24, True, BLACK, PP_ALIGN.CENTER)
    add_text(slide, "Estimation et dynamique des distributions risque-neutres\nà partir d’options SPX", 1.1, 1.65, 11.1, 1.35, 30, True, BLACK, PP_ALIGN.CENTER)
    add_text(slide, "Réunion de suivi n°1  ·  23 octobre 2026", 1.0, 3.35, 11.3, 0.45, 17, False, BLUE, PP_ALIGN.CENTER)
    add_text(slide, "Rayan Ait Arkoub  ·  Matthieu Balliste  ·  Ilan Chahdi  ·  Nathan Denoux  ·  Jordi Paris", 0.8, 5.25, 11.7, 0.55, 15, False, DARK, PP_ALIGN.CENTER)
    add_text(slide, "Tuteur : Matthieu Garcin", 1.0, 5.78, 11.3, 0.4, 13, False, MUTED, PP_ALIGN.CENTER)

    # 2 — Question et protocole
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "1. Question de recherche et protocole", 2, prs)
    add_text(slide, "Peut-on prévoir une densité risque-neutre suffisamment stable pour améliorer la valorisation ou la couverture d’options ?", 0.9, 1.28, 11.6, 0.85, 23, True, BLACK, PP_ALIGN.CENTER)
    columns = [
        ("Estimation", "Cotations SPX → smile commune → RND", "Breeden–Litzenberger\nCOS\nRepondération entropique"),
        ("Dynamique", "RND quotidiennes à maturité constante", "Distances entre densités\nFacteurs fonctionnels\nPrévision AR/VAR"),
        ("Évaluation", "Prévisions hors échantillon", "Erreur de pricing\nErreur de couverture\nCoûts de transaction"),
    ]
    for i, (head, sub, body) in enumerate(columns):
        x = 0.75 + i * 4.18
        add_text(slide, head, x, 2.55, 3.7, 0.38, 19, True, BLACK, PP_ALIGN.CENTER)
        add_text(slide, sub, x, 3.05, 3.7, 0.6, 14, True, BLUE, PP_ALIGN.CENTER)
        add_text(slide, body, x, 3.85, 3.7, 1.5, 15, False, DARK, PP_ALIGN.CENTER)
    add_text(slide, "Jalon actuel : pipeline reproductible sur une date et diagnostic de la structure par terme.", 1.0, 6.20, 11.3, 0.42, 16, True, BLACK, PP_ALIGN.CENTER)

    # 3 — Données
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "2. Données Cboe et règles de filtrage", 3, prs)
    add_picture_contain(slide, RESULTS / "01_data_quality.png", 0.55, 1.15, 7.75, 5.35)
    add_metric_line(slide, "1 446", "contrats SPX", 8.62, 1.28)
    add_metric_line(slide, "1 277", "cotations valides", 8.62, 1.83)
    add_metric_line(slide, "338", "points de calibration", 8.62, 2.38)
    add_metric_line(slide, "91 j", "échéance du 15/01/2027", 8.62, 2.93)
    add_bullets(slide, [
        "Mid bid–ask plutôt que dernier prix parfois ancien.",
        "Exclusion : bid nul, marché croisé, mid < 0,05 et spread relatif > 35 %.",
        "Cboe gratuit : prototype ponctuel ; Bloomberg : historique et stratégie.",
    ], 8.55, 3.65, 4.2, 2.45, 14)
    add_caption(slide, "Snapshot Cboe différé du 16/10/2026 à 17:12:17.", 0.8, 6.52, 7.2)

    # 4 — Cadre théorique
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "3. Identités théoriques utilisées", 4, prs)
    add_text(slide, "Prix du call sous la mesure risque-neutre", 0.8, 1.18, 11.8, 0.35, 16, True, BLACK)
    add_picture_contain(slide, assets["call"], 1.5, 1.52, 10.3, 0.72)
    add_text(slide, "Estimation conjointe du forward et de l’actualisation", 0.8, 2.55, 11.8, 0.35, 16, True, BLACK)
    add_picture_contain(slide, assets["parity"], 1.5, 2.88, 10.3, 0.72)
    add_text(slide, "Extraction de la densité", 0.8, 3.93, 11.8, 0.35, 16, True, BLACK)
    add_picture_contain(slide, assets["bl"], 1.5, 4.26, 10.3, 0.78)
    add_bullets(slide, [
        "Options européennes, même sous-jacent et même échéance.",
        "La convexité de C(K,T) garantit q(K,T) ≥ 0 en absence d’arbitrage.",
        "La double dérivation impose de lisser avant de dériver.",
    ], 1.0, 5.40, 11.2, 1.2, 14)

    # 5 — SVI
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "4. Une smile commune : calibration SVI", 5, prs)
    add_picture_contain(slide, RESULTS / "02_svi_smile.png", 0.45, 1.10, 7.2, 4.75)
    add_picture_contain(slide, assets["svi"], 7.55, 1.27, 5.35, 0.95)
    svi = DIAG["svi_parameters"]
    add_bullets(slide, [
        "Un seul vecteur de cinq paramètres pour les 338 options.",
        f"F = {DIAG['forward_from_parity']:,.2f} ; D = {DIAG['discount_factor_from_parity']:.5f} ; r = {100*DIAG['implied_continuous_rate']:.2f} %.",
        f"RMSE de parité = {DIAG['put_call_parity_rmse']:.3f} point.",
        f"RMSE de reconstruction = {DIAG['call_price_rmse']:.3f} point.",
        f"ρ = {svi['rho']:.3f} ; σ = {svi['sigma']:.3f}.",
    ], 7.75, 2.55, 4.85, 3.35, 15)
    add_text(slide, "SVI lisse la coupe ; Breeden–Litzenberger reste la méthode d’extraction.", 0.9, 6.38, 11.6, 0.38, 15, True, BLACK, PP_ALIGN.CENTER)

    # 6 — Courbe puis densité
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "5. De la courbe de calls à la RND", 6, prs)
    add_picture_contain(slide, RESULTS / "03_call_curve.png", 0.35, 1.12, 6.2, 4.55)
    add_picture_contain(slide, RESULTS / "04_risk_neutral_density.png", 6.72, 1.12, 6.2, 4.55)
    add_caption(slide, "Courbe de calls reconstruite sur 1 601 strikes.", 0.55, 5.70, 5.8)
    add_caption(slide, "Densité obtenue par dérivée seconde de la courbe lissée.", 6.85, 5.70, 5.8)
    add_picture_contain(slide, assets["bl"], 2.2, 6.07, 8.9, 0.58)

    # 7 — Diagnostics 91 jours
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "6. Diagnostics de la RND à 91 jours", 7, prs)
    add_picture_contain(slide, assets["support_diag"], 0.45, 1.10, 8.5, 4.95)
    gap = DIAG["normalized_density_mean"] - DIAG["forward_from_parity"]
    add_metric_line(slide, f"{DIAG['raw_density_mass_on_grid']:.5f}", "masse brute", 9.18, 1.35, 3.65)
    add_metric_line(slide, f"{DIAG['absolute_negative_mass']:.6f}", "masse négative", 9.18, 2.02, 3.65)
    add_metric_line(slide, f"+{gap:.2f}", "moyenne − forward", 9.18, 2.69, 3.65)
    add_bullets(slide, [
        "Aucune négativité matérielle au-delà de 0,1 % du pic.",
        "0,47 % de masse manque principalement dans la queue gauche.",
        "Le support fini interdit d’interpréter les queues sans test de robustesse.",
    ], 9.08, 3.55, 3.9, 2.35, 13)
    add_text(slide, "Diagnostic, pas correction définitive : eSSVI ou lissage convexe à comparer.", 0.75, 6.35, 12.0, 0.38, 15, True, RED, PP_ALIGN.CENTER)

    # 8 — Surface 3D
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "7. Structure par terme : six maturités, une seule date", 8, prs)
    add_picture_contain(slide, RESULTS / "05_rnd_surface_3d_maturity.png", 0.35, 1.02, 9.0, 5.92)
    add_picture_contain(slide, assets["change"], 9.25, 1.33, 3.75, 0.72)
    add_bullets(slide, [
        "31, 63, 91, 126, 154 et 182 jours.",
        "Grille commune de log-moneyness.",
        "Élargissement de la densité avec l’horizon.",
        "Lignes noires : coupes calibrées ; surface : interpolation visuelle PCHIP.",
        "Ce n’est pas encore une dynamique temporelle.",
    ], 9.35, 2.35, 3.55, 3.9, 13)

    # 9 — Fiabilité selon maturité
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "8. Toutes les maturités ne sont pas également fiables", 9, prs)
    add_picture_contain(slide, RESULTS / "06_rnd_heatmap_maturity.png", 0.35, 1.15, 6.1, 4.9)
    add_picture_contain(slide, assets["term_diag"], 6.45, 1.15, 6.55, 4.9)
    add_text(slide, "La coupe 31 jours est conservée pour le diagnostic, mais exclue des conclusions : masse 1,138 et moyenne 5,4 % au-dessus du forward.", 0.75, 6.24, 11.85, 0.55, 14, True, RED, PP_ALIGN.CENTER)

    # 10 — Dynamique et stratégie
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "9. Passage à la dynamique et test de stratégie", 10, prs)
    add_text(slide, "Réduction fonctionnelle", 0.8, 1.20, 5.8, 0.35, 17, True, BLACK, PP_ALIGN.CENTER)
    add_picture_contain(slide, assets["factor"], 0.55, 1.62, 6.1, 0.82)
    add_text(slide, "Prévision des scores", 6.75, 1.20, 5.8, 0.35, 17, True, BLACK, PP_ALIGN.CENTER)
    add_picture_contain(slide, assets["var"], 6.55, 1.62, 6.2, 0.82)
    add_bullets(slide, [
        "Une RND quotidienne à maturité constante, extraite de l’historique Bloomberg.",
        "Comparaison : marche aléatoire fonctionnelle, AR et VAR.",
        "Distances : Jensen–Shannon, Wasserstein et erreur de pricing.",
    ], 0.85, 3.0, 5.75, 2.35, 15)
    add_bullets(slide, [
        "Couverture RND courante vs RND prévue vs Black–Scholes.",
        "Évaluation hors échantillon : erreur de couverture, variance, drawdown.",
        "Résultats nets des coûts de transaction.",
    ], 6.78, 3.0, 5.7, 2.35, 15)
    add_text(slide, "Cboe gratuit pour le prototype ; Bloomberg pour l’historique et le backtest.", 1.1, 6.12, 11.1, 0.45, 17, True, BLACK, PP_ALIGN.CENTER)

    # 11 — Conclusion
    slide = prs.slides.add_slide(blank); set_background(slide, prs); add_header(slide, "10. Résultats acquis et prochain jalon", 11, prs)
    add_text(slide, "Acquis", 0.9, 1.25, 5.5, 0.4, 20, True, BLACK)
    add_bullets(slide, [
        "Chaîne SPX réelle nettoyée et documentée.",
        "Forward, actualisation et SVI estimés sur 338 points.",
        "RND Breeden–Litzenberger contrôlée à 91 jours.",
        "Structure par terme et limites identifiées.",
        "Python–Julia : écart moyen 3,4 × 10⁻⁹ ; maximum 2,2 × 10⁻⁶.",
    ], 0.9, 1.88, 5.6, 3.85, 16)
    add_text(slide, "Prochain jalon", 6.85, 1.25, 5.5, 0.4, 20, True, BLACK)
    add_bullets(slide, [
        "Sécuriser l’export historique Bloomberg.",
        "Construire les RND à maturité constante.",
        "Comparer BL, COS et repondération entropique.",
        "Estimer les facteurs puis réaliser le backtest.",
    ], 6.85, 1.88, 5.55, 3.25, 16)
    add_text(slide, "Décision attendue : horizon constant retenu et protocole d’accès Bloomberg.", 1.0, 6.10, 11.25, 0.55, 18, True, BLUE, PP_ALIGN.CENTER)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Support_suivi_23-10-2026.pptx"
    prs.save(path)
    return path


def export_pdf_with_powerpoint(pptx: Path) -> Path:
    pdf = pptx.with_suffix(".pdf")
    q_pptx = str(pptx.resolve()).replace("'", "''")
    q_pdf = str(pdf.resolve()).replace("'", "''")
    script = (
        "$ErrorActionPreference='Stop'; "
        "$app=New-Object -ComObject PowerPoint.Application; "
        f"$deck=$app.Presentations.Open('{q_pptx}', $true, $false, $false); "
        f"$deck.SaveAs('{q_pdf}', 32); "
        "$deck.Close(); $app.Quit();"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return pdf


def build_speaker_script_pdf() -> Path:
    """Crée une version imprimable des notes orales Markdown."""
    source = ROOT / "Notes_orales_suivi_23-10-2026.md"
    destination = OUT / "Script_oral_suivi_23-10-2026.pdf"
    pdfmetrics.registerFont(TTFont("ArialLocal", r"C:\Windows\Fonts\arial.ttf"))
    pdfmetrics.registerFont(TTFont("ArialLocal-Bold", r"C:\Windows\Fonts\arialbd.ttf"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleLocal", parent=styles["Title"], fontName="ArialLocal-Bold", fontSize=18, leading=22, alignment=TA_CENTER, textColor=rl_colors.HexColor("#" + BLACK), spaceAfter=12)
    h1 = ParagraphStyle("H1Local", parent=styles["Heading1"], fontName="ArialLocal-Bold", fontSize=13, leading=16, textColor=rl_colors.HexColor("#" + BLACK), spaceBefore=10, spaceAfter=5, keepWithNext=True)
    h2 = ParagraphStyle("H2Local", parent=styles["Heading2"], fontName="ArialLocal-Bold", fontSize=10.5, leading=13, textColor=rl_colors.HexColor("#" + BLUE), spaceBefore=7, spaceAfter=3, keepWithNext=True)
    body = ParagraphStyle("BodyLocal", parent=styles["BodyText"], fontName="ArialLocal", fontSize=9.2, leading=12.5, alignment=TA_JUSTIFY, textColor=rl_colors.HexColor("#" + DARK), spaceAfter=6)

    def markup(value: str) -> str:
        value = html.escape(value)
        value = value.replace("$3{,}4\\times10^{-9}$", "3,4 × 10<super>−9</super>")
        value = value.replace("$2{,}2\\times10^{-6}$", "2,2 × 10<super>−6</super>")
        value = value.replace("$10^{-8}$", "10<super>−8</super>")
        value = value.replace("$10^{-3}$", "10<super>−3</super>")
        value = value.replace("$F_{t,T}=\\mathbb{E}^{\\mathbb{Q}}_t[S_T]$", "F(t,T) = Eᴼ[S(T)]")
        return value

    story = []
    paragraph_lines: list[str] = []

    def flush() -> None:
        if paragraph_lines:
            story.append(Paragraph(markup(" ".join(paragraph_lines)), body))
            paragraph_lines.clear()

    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            flush()
        elif line.startswith("# "):
            flush(); story.append(Paragraph(markup(line[2:]), title))
        elif line.startswith("## "):
            flush(); story.append(Paragraph(markup(line[3:]), h1))
        elif line.startswith("### "):
            flush(); story.append(Paragraph(markup(line[4:]), h2))
        else:
            paragraph_lines.append(line)
    flush()

    def footer(canvas, doc) -> None:
        canvas.saveState(); canvas.setFont("ArialLocal", 8); canvas.setFillColorRGB(0.35, 0.39, 0.42)
        canvas.drawString(2.0 * cm, 1.0 * cm, "Projet PI² n°176 — Script oral — 23 octobre 2026")
        canvas.drawRightString(A4[0] - 2.0 * cm, 1.0 * cm, str(doc.page)); canvas.restoreState()

    doc = SimpleDocTemplate(str(destination), pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=1.8 * cm, bottomMargin=1.7 * cm)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    shutil.copy2(source, OUT / source.name)
    return destination


if __name__ == "__main__":
    pptx_path = build_presentation()
    pdf_path = export_pdf_with_powerpoint(pptx_path)
    script_path = build_speaker_script_pdf()
    print(pptx_path)
    print(pdf_path)
    print(script_path)
